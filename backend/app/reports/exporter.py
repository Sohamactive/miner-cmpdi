"""DOCX export — cover page, headings, fact tables, chart placeholders, citations."""

from __future__ import annotations

import json
import logging
import os
import re
from pathlib import Path

from app.extraction.postgres import create_session
from app.extraction.models import ClaimRecord, Report

logger = logging.getLogger(__name__)

EXPORT_FOLDER = os.getenv("EXPORT_FOLDER", "data/reports")


def _ensure_export_dir() -> Path:
    p = Path(EXPORT_FOLDER)
    p.mkdir(parents=True, exist_ok=True)
    return p


def _clean_report_text(value: object) -> str:
    text = str(value or "")
    text = re.sub(r"\s*\[[FC]\d+(?:,\s*doc=[^\]]+)?\]", "", text)
    return re.sub(r"\s+", " ", text).strip()


def _shade_cell(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _style_table(table, *, header_fill: str = "1F4E78") -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.shared import Pt, RGBColor

    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        if row_index == 0:
            for cell in row.cells:
                _shade_cell(cell, header_fill)


def _add_claim_table(doc, claims) -> None:
    rows = [
        ["Claim", "Value", "Unit", "Document", "Page", "Status"],
    ]
    for claim in claims:
        if claim.validation_status == "UNSUPPORTED":
            continue
        value = "N/A" if claim.value is None else f"{claim.value:g}"
        rows.append([
            _clean_report_text(claim.claim_text)[:180],
            value,
            claim.unit or "N/A",
            claim.document_id or "N/A",
            claim.page_range or "N/A",
            claim.validation_status,
        ])

    if len(rows) == 1:
        doc.add_paragraph("No validated claims were extracted.")
        return

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row_data in enumerate(rows):
        for column_index, cell_text in enumerate(row_data):
            table.cell(row_index, column_index).text = str(cell_text)
    _style_table(table)
    doc.add_paragraph("")


def _add_citations(doc, claims) -> None:
    seen: set[str] = set()
    citations: list[tuple[str, str]] = []
    for claim in claims:
        if claim.validation_status == "UNSUPPORTED":
            continue
        key = f"{claim.document_id}:{claim.page_range}"
        if key in seen:
            continue
        seen.add(key)
        location = f"Document {claim.document_id or 'N/A'}, Page {claim.page_range or 'N/A'}"
        citations.append((location, _clean_report_text(claim.evidence_text)[:300]))

    if not citations:
        doc.add_paragraph("No supported citations were available.")
        return
    for index, (location, excerpt) in enumerate(citations, start=1):
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(f"{location}: ").bold = True
        paragraph.add_run(excerpt)


def _add_review_summary(doc, claims) -> None:
    counts = {
        "SUPPORTED": sum(c.validation_status == "SUPPORTED" for c in claims),
        "NEEDS_REVIEW": sum(c.validation_status == "NEEDS_REVIEW" for c in claims),
        "UNSUPPORTED": sum(c.validation_status == "UNSUPPORTED" for c in claims),
    }
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Validation status"
    table.cell(0, 1).text = "Claims"
    for label, count in counts.items():
        row = table.add_row()
        row.cells[0].text = label.replace("_", " ").title()
        row.cells[1].text = str(count)
    _style_table(table)


def _add_toc_field(doc) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "TOC \\o \"1-3\" \\h \\z \\u"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click to update this table of contents."
    separate.append(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def _add_overview(doc, report, sections, claims) -> None:
    from docx.shared import Pt

    overview = next(
        (section for section in sections if section.get("heading") in {"preamble", "overview"}),
        None,
    )
    overview_text = (overview or {}).get("paragraphs", [])
    if overview_text:
        paragraph = doc.add_paragraph(_clean_report_text(overview_text[0]))
        paragraph.paragraph_format.space_after = Pt(8)
    else:
        doc.add_paragraph(
            "This evidence-grounded report summarizes the available indexed records "
            "and identifies the supporting documents and page references."
        )

    summary = next(
        (section for section in sections if section.get("heading") in {"facts_summary", "analysis"}),
        None,
    )
    if summary and summary.get("paragraphs"):
        doc.add_paragraph("Key findings", style="Heading 2")
        for paragraph_text in summary["paragraphs"][:3]:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(_clean_report_text(paragraph_text))

    evidence_count = sum(
        1 for claim in claims if claim.document_id or claim.evidence_text
    )
    supported_count = sum(
        claim.validation_status == "SUPPORTED" for claim in claims
    )
    metadata = doc.add_table(rows=4, cols=2)
    overview_rows = [
        ("Report status", report.status.replace("_", " ").title()),
        ("Evidence-backed claims", str(evidence_count)),
        ("Supported claims", str(supported_count)),
        ("Version", str(report.version)),
    ]
    for row, (label, value) in zip(metadata.rows, overview_rows, strict=True):
        row.cells[0].text = label
        row.cells[1].text = value
    _style_table(metadata, header_fill="5B7C99")
    doc.add_paragraph("")


def export_docx(*, report_id: int) -> str:
    """Export a report to DOCX. Returns file path."""
    try:
        from docx import Document as WordDocument
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: uv add python-docx")

    session = create_session()
    try:
        report = session.get(Report, report_id)
        if report is None:
            raise ValueError(f"Report {report_id} not found")

        claims = list(session.scalars(
            __import__("sqlalchemy").select(ClaimRecord).where(ClaimRecord.report_id == report_id)
        ))
    finally:
        session.close()

    sections = json.loads(report.sections_json or "[]")
    export_dir = _ensure_export_dir()
    filepath = str(export_dir / f"report-{report_id}.docx")

    doc = WordDocument()

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings._element.append(update_fields)

    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    doc.styles["Title"].font.name = "Aptos Display"
    doc.styles["Title"].font.color.rgb = RGBColor(31, 78, 120)
    for heading_name in ("Heading 1", "Heading 2"):
        heading_style = doc.styles[heading_name]
        heading_style.font.name = "Aptos Display"
        heading_style.font.color.rgb = RGBColor(31, 78, 120)

    header = section.header.paragraphs[0]
    header.text = "M.I.N.E.R. | Evidence-grounded report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Confidential working draft | Page ")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    # Cover page
    style = doc.styles["Title"]
    style.font.size = Pt(24)
    title = doc.add_paragraph(report.title, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run(f"{report.report_type.replace('_', ' ').title()}\n").bold = True
    metadata.add_run(f"Version {report.version}")
    doc.add_paragraph("")
    doc.add_heading("Report Overview", level=1)
    _add_overview(doc, report, sections, claims)
    doc.add_page_break()

    # Word-updatable table of contents
    doc.add_heading("Table of Contents", level=1)
    _add_toc_field(doc)
    doc.add_page_break()

    # Sections
    for section_data in sections:
        heading = section_data.get("heading", "Untitled")
        paragraphs = section_data.get("paragraphs", [])

        if heading.lower() == "citations":
            continue

        doc.add_heading(heading.replace("_", " ").title(), level=1)

        for para_text in paragraphs:
            paragraph = doc.add_paragraph(_clean_report_text(para_text))
            paragraph.paragraph_format.space_after = Pt(6)

        if heading in {"evidence_table", "production_figures"}:
            _add_claim_table(doc, claims)

        # Chart placeholder
        chart_spec = section_data.get("chart_spec")
        if chart_spec:
            chart_note = doc.add_paragraph()
            chart_note.add_run("Chart specification: ").bold = True
            chart_note.add_run(
                f"{chart_spec.get('type', 'bar')} chart of "
                f"{chart_spec.get('x_metric', '')} versus {chart_spec.get('y_metric', '')}."
            )
            doc.add_paragraph("")

    # Citations section
    doc.add_heading("Citations", level=1)
    _add_citations(doc, claims)

    # Review summary
    doc.add_heading("Review Summary", level=1)
    _add_review_summary(doc, claims)

    doc.save(filepath)
    logger.info("DOCX exported: %s (report_id=%d)", filepath, report_id)
    return filepath
